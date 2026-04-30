from __future__ import annotations

import json
import math
from collections import Counter, defaultdict
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

INPUT = Path('/home/ubuntu/upload/INFOGERAL.xlsx')
OUT_DIR = Path('/home/ubuntu/infogeral-dashboard-local/analysis')
OUT_DIR.mkdir(parents=True, exist_ok=True)


def clean_value(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, (datetime, date, time)):
        return v.isoformat()
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return None
        if v.is_integer():
            return int(v)
    return v


def as_text(v: Any) -> str:
    v = clean_value(v)
    if v is None:
        return ''
    return str(v).strip()


def detect_header_row(rows: list[list[Any]], max_scan: int = 20) -> int:
    best_idx = 0
    best_score = -1
    for idx, row in enumerate(rows[:max_scan]):
        non_empty = [as_text(c) for c in row if as_text(c)]
        unique = len(set(non_empty))
        score = len(non_empty) + unique * 0.35
        # Penaliza linhas muito numéricas, favorece rótulos de coluna.
        numeric_like = sum(1 for x in non_empty if x.replace('.', '', 1).replace(',', '', 1).isdigit())
        score -= numeric_like * 0.4
        if score > best_score:
            best_idx = idx
            best_score = score
    return best_idx


def normalize_headers(header: list[Any]) -> list[str]:
    seen: Counter[str] = Counter()
    headers: list[str] = []
    for i, h in enumerate(header, 1):
        name = as_text(h) or f'COLUNA_{i}'
        name = ' '.join(name.split())
        seen[name] += 1
        if seen[name] > 1:
            name = f'{name}_{seen[name]}'
        headers.append(name)
    return headers


def row_to_record(headers: list[str], row: list[Any]) -> dict[str, Any]:
    return {headers[i]: clean_value(row[i]) if i < len(row) else None for i in range(len(headers))}


def is_blank_record(record: dict[str, Any]) -> bool:
    return all(as_text(v) == '' for v in record.values())


wb = load_workbook(INPUT, data_only=False)
wb_values = load_workbook(INPUT, data_only=True)
summary: dict[str, Any] = {
    'arquivo': str(INPUT),
    'abas': [],
    'total_abas': len(wb.sheetnames),
}
all_data: dict[str, Any] = {}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    wsv = wb_values[sheet_name]
    rows_formula = [list(r) for r in ws.iter_rows(values_only=True)]
    rows_values = [list(r) for r in wsv.iter_rows(values_only=True)]
    header_idx = detect_header_row(rows_values) if rows_values else 0
    header = rows_values[header_idx] if rows_values else []
    headers = normalize_headers(header)
    records = []
    for row in rows_values[header_idx + 1:]:
        padded = list(row) + [None] * max(0, len(headers) - len(row))
        record = row_to_record(headers, padded[:len(headers)])
        if not is_blank_record(record):
            records.append(record)

    formula_cells = []
    for row in rows_formula:
        for cell in row:
            if isinstance(cell, str) and cell.startswith('='):
                formula_cells.append(cell)

    column_profiles = []
    for h in headers:
        values = [r.get(h) for r in records]
        non_empty = [clean_value(v) for v in values if as_text(v) != '']
        unique_values = list(dict.fromkeys(non_empty))[:25]
        type_counter = Counter(type(v).__name__ for v in non_empty)
        text_counter = Counter(as_text(v) for v in non_empty)
        column_profiles.append({
            'campo': h,
            'preenchidos': len(non_empty),
            'vazios': len(records) - len(non_empty),
            'unicos': len(set(as_text(v) for v in non_empty)),
            'tipos': dict(type_counter),
            'amostras': unique_values,
            'valores_mais_frequentes': text_counter.most_common(10),
        })

    sheet_summary = {
        'aba': sheet_name,
        'linhas_excel': ws.max_row,
        'colunas_excel': ws.max_column,
        'linha_cabecalho_detectada': header_idx + 1,
        'campos': headers,
        'total_registros_detectados': len(records),
        'celulas_mescladas': [str(rng) for rng in ws.merged_cells.ranges],
        'quantidade_formulas': len(formula_cells),
        'amostras_formulas': formula_cells[:20],
        'perfil_campos': column_profiles,
        'primeiros_registros': records[:5],
    }
    summary['abas'].append(sheet_summary)
    all_data[sheet_name] = {
        'headers': headers,
        'records': records,
        'summary': sheet_summary,
    }

(OUT_DIR / 'workbook_summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
(OUT_DIR / 'workbook_data.json').write_text(json.dumps(all_data, ensure_ascii=False, indent=2), encoding='utf-8')

# Markdown de leitura humana
md = ['# Análise estrutural do arquivo INFOGERAL.xlsx', '']
md.append(f"Arquivo analisado: `{INPUT}`")
md.append(f"Total de abas: **{summary['total_abas']}**")
md.append('')
for aba in summary['abas']:
    md.append(f"## Aba: {aba['aba']}")
    md.append('')
    md.append(f"Linha de cabeçalho detectada: **{aba['linha_cabecalho_detectada']}**")
    md.append(f"Dimensão Excel: **{aba['linhas_excel']} linhas x {aba['colunas_excel']} colunas**")
    md.append(f"Registros detectados: **{aba['total_registros_detectados']}**")
    md.append(f"Fórmulas detectadas: **{aba['quantidade_formulas']}**")
    md.append('')
    md.append('| Campo | Preenchidos | Vazios | Únicos | Tipos | Valores frequentes |')
    md.append('| --- | ---: | ---: | ---: | --- | --- |')
    for col in aba['perfil_campos']:
        freq = ', '.join([f'{k} ({v})' for k, v in col['valores_mais_frequentes'][:5]])
        md.append(f"| {col['campo']} | {col['preenchidos']} | {col['vazios']} | {col['unicos']} | {json.dumps(col['tipos'], ensure_ascii=False)} | {freq} |")
    md.append('')
    md.append('### Primeiros registros')
    md.append('')
    if aba['primeiros_registros']:
        fields = aba['campos'][:min(8, len(aba['campos']))]
        md.append('| ' + ' | '.join(fields) + ' |')
        md.append('| ' + ' | '.join(['---'] * len(fields)) + ' |')
        for rec in aba['primeiros_registros']:
            md.append('| ' + ' | '.join(as_text(rec.get(f)).replace('|', '\\|') for f in fields) + ' |')
    else:
        md.append('Sem registros detectados.')
    md.append('')

(OUT_DIR / 'analise_infogeral.md').write_text('\n'.join(md), encoding='utf-8')
print(json.dumps({'summary': str(OUT_DIR / 'workbook_summary.json'), 'data': str(OUT_DIR / 'workbook_data.json'), 'markdown': str(OUT_DIR / 'analise_infogeral.md')}, ensure_ascii=False, indent=2))
