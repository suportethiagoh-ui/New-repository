from __future__ import annotations

import argparse
import json
import math
import re
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parent
DEFAULT_EXCEL = Path('/home/ubuntu/upload/INFOGERAL.xlsx')
DEFAULT_OUT = PROJECT / 'client' / 'src' / 'data' / 'infogeralData.ts'
RAW_JSON_OUT = PROJECT / 'analysis' / 'infogeral_normalizado.json'

SETOR_SHEETS = {
    'Fundição': 'Fundição',
    'Classificação': 'Classificação',
    'Inspção P. Pequenas': 'Inspeção Peças Pequenas',
    'Esmaltação': 'Esmaltação',
}


def clean_value(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, (datetime, date, time)):
        return v.isoformat()
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return None
        if v.is_integer():
            return int(v)
    if isinstance(v, str):
        s = ' '.join(v.replace('\u00a0', ' ').split())
        return s if s else None
    return v


def text(v: Any) -> str:
    v = clean_value(v)
    return '' if v is None else str(v).strip()


def parse_int_like(v: Any) -> int | None:
    s = text(v)
    if not s:
        return None
    m = re.search(r'\d+', s)
    if not m:
        return None
    return int(m.group(0))


def parse_rows(ws, start_row: int, columns: int = 5) -> list[dict[str, Any]]:
    rows = []
    for r in range(start_row, ws.max_row + 1):
        vals = [clean_value(ws.cell(r, c).value) for c in range(1, columns + 1)]
        if not any(text(v) for v in vals):
            continue
        first = text(vals[0]).lower()
        if first in {'nomes/cod', 'nomes/cod.', 'nomes/codigo', 'novato'}:
            continue
        rows.append({
            'tipo': vals[0],
            'codigo': vals[1],
            'carimbo': vals[2],
            'nome': vals[3],
            'nomeCompleto': vals[4],
            'linhaExcel': r,
        })
    return rows


def extract_raw_sheet(ws) -> dict[str, Any]:
    rows = []
    for r in range(1, ws.max_row + 1):
        vals = [clean_value(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if any(text(v) for v in vals):
            rows.append({'linhaExcel': r, 'valores': vals})
    return {
        'nome': ws.title,
        'linhas': ws.max_row,
        'colunas': ws.max_column,
        'linhasComDados': rows,
        'celulasMescladas': [str(rng) for rng in ws.merged_cells.ranges],
    }


def build_payload(excel_path: Path) -> dict[str, Any]:
    wb = load_workbook(excel_path, data_only=True)

    colaboradores: list[dict[str, Any]] = []
    if 'Geral Cod' in wb.sheetnames:
        ws = wb['Geral Cod']
        for r in range(2, ws.max_row + 1):
            codigo = clean_value(ws.cell(r, 1).value)
            carimbo = clean_value(ws.cell(r, 2).value)
            nome = clean_value(ws.cell(r, 3).value)
            nome_completo = clean_value(ws.cell(r, 4).value)
            setor = clean_value(ws.cell(r, 5).value)
            if not any(text(v) for v in [codigo, carimbo, nome, nome_completo, setor]):
                continue
            colaboradores.append({
                'id': f'colab-{len(colaboradores)+1:03d}',
                'codigo': codigo,
                'carimbo': carimbo,
                'carimboNumero': parse_int_like(carimbo),
                'nome': nome,
                'nomeCompleto': nome_completo,
                'setor': setor,
                'linhaExcel': r,
                'origem': 'Geral Cod',
            })

    detalhes_por_setor: dict[str, list[dict[str, Any]]] = {}
    for sheet_name, setor in SETOR_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            detalhes_por_setor[setor] = []
            continue
        ws = wb[sheet_name]
        registros = parse_rows(ws, start_row=1, columns=5)
        for idx, item in enumerate(registros, 1):
            item['id'] = f'{setor.lower().replace(" ", "-")}-{idx:03d}'
            item['setor'] = setor
            item['origem'] = sheet_name
            item['carimboNumero'] = parse_int_like(item.get('carimbo'))
        detalhes_por_setor[setor] = registros

    pecas = []
    if 'PEÇAS' in wb.sheetnames:
        ws = wb['PEÇAS']
        for r in range(1, ws.max_row + 1):
            nome = clean_value(ws.cell(r, 1).value)
            codigo = clean_value(ws.cell(r, 2).value)
            if not text(nome) or text(nome).upper() == 'PEÇAS':
                continue
            pecas.append({
                'id': f'peca-{len(pecas)+1:03d}',
                'nome': nome,
                'codigo': codigo,
                'linhaExcel': r,
                'origem': 'PEÇAS',
            })

    raw_sheets = {name: extract_raw_sheet(wb[name]) for name in wb.sheetnames}

    setores = Counter(text(c['setor']) for c in colaboradores if text(c.get('setor')))
    tipos = Counter(text(c.get('tipo')) for registros in detalhes_por_setor.values() for c in registros if text(c.get('tipo')))
    carimbos = [c.get('carimboNumero') for c in colaboradores if isinstance(c.get('carimboNumero'), int)]
    nomes_duplicados = [nome for nome, count in Counter(text(c.get('nome')) for c in colaboradores).items() if nome and count > 1]

    payload = {
        'metadata': {
            'arquivoOrigem': str(excel_path),
            'geradoEm': datetime.now().isoformat(timespec='seconds'),
            'totalAbas': len(wb.sheetnames),
            'abas': wb.sheetnames,
            'totalColaboradores': len(colaboradores),
            'totalPecas': len(pecas),
            'setores': dict(setores),
            'tipos': dict(tipos),
            'menorCarimboNumerico': min(carimbos) if carimbos else None,
            'maiorCarimboNumerico': max(carimbos) if carimbos else None,
            'nomesDuplicados': nomes_duplicados,
            'observacao': 'A aba Geral Cod foi usada como cadastro mestre. As abas setoriais foram preservadas como detalhamento operacional e todas as abas brutas foram mantidas em rawSheets.',
        },
        'colaboradores': colaboradores,
        'detalhesPorSetor': detalhes_por_setor,
        'pecas': pecas,
        'rawSheets': raw_sheets,
    }
    return payload


def write_ts(payload: dict[str, Any], out: Path) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    json_payload = json.dumps(payload, ensure_ascii=False, indent=2)
    content = f"""// Filosofia visual aplicada neste arquivo: Swiss Corporate Brutalism editorial.
// Este arquivo é gerado automaticamente por carregar_infogeral.py a partir do Excel INFOGERAL.
// Não edite manualmente os dados abaixo; rode novamente: python3 carregar_infogeral.py --excel /caminho/INFOGERAL.xlsx

export type Colaborador = {{
  id: string;
  codigo: string | number | null;
  carimbo: string | number | null;
  carimboNumero: number | null;
  nome: string | null;
  nomeCompleto: string | null;
  setor: string | null;
  linhaExcel: number;
  origem: string;
}};

export type Peca = {{
  id: string;
  nome: string | null;
  codigo: string | number | null;
  linhaExcel: number;
  origem: string;
}};

export const infogeralData = {json_payload} as const;

export const colaboradores = infogeralData.colaboradores as readonly Colaborador[];
export const pecas = infogeralData.pecas as readonly Peca[];
export const detalhesPorSetor = infogeralData.detalhesPorSetor;
export const rawSheets = infogeralData.rawSheets;
export const metadata = infogeralData.metadata;
"""
    out.write_text(content, encoding='utf-8')


def main() -> None:
    parser = argparse.ArgumentParser(description='Carrega a planilha INFOGERAL.xlsx e gera os dados locais da aplicação.')
    parser.add_argument('--excel', default=str(DEFAULT_EXCEL), help='Caminho do arquivo .xlsx de entrada.')
    parser.add_argument('--saida', default=str(DEFAULT_OUT), help='Arquivo TypeScript de saída.')
    parser.add_argument('--json', default=str(RAW_JSON_OUT), help='Arquivo JSON normalizado para auditoria.')
    args = parser.parse_args()

    excel_path = Path(args.excel).expanduser().resolve()
    payload = build_payload(excel_path)
    write_ts(payload, Path(args.saida).expanduser().resolve())
    json_out = Path(args.json).expanduser().resolve()
    json_out.parent.mkdir(parents=True, exist_ok=True)
    json_out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')

    print(json.dumps({
        'status': 'ok',
        'excel': str(excel_path),
        'saida_ts': str(Path(args.saida).expanduser().resolve()),
        'saida_json': str(json_out),
        'total_colaboradores': payload['metadata']['totalColaboradores'],
        'total_pecas': payload['metadata']['totalPecas'],
        'setores': payload['metadata']['setores'],
        'abas': payload['metadata']['abas'],
    }, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
