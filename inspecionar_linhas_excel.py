from pathlib import Path
from openpyxl import load_workbook

INPUT = Path('/home/ubuntu/upload/INFOGERAL.xlsx')
wb = load_workbook(INPUT, data_only=True)
for sheet in wb.sheetnames:
    ws = wb[sheet]
    print('\n' + '='*90)
    print(f'ABA: {sheet} | max_row={ws.max_row} | max_col={ws.max_column}')
    print('='*90)
    for r in range(1, min(ws.max_row, 25) + 1):
        vals = []
        for c in range(1, min(ws.max_column, 12) + 1):
            v = ws.cell(r, c).value
            vals.append('' if v is None else str(v))
        print(f'{r:02d}: ' + ' | '.join(vals))
